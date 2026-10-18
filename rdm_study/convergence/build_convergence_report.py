"""Research supplement from the completed, verified convergence experiment."""
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import argparse,json
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from common import ROOT,CORE,LONG
from labels import LABELS

HERE=Path(__file__).resolve().parent
OUT=ROOT.parent/'deliverables/Strategy_RDM_Sampling_Convergence_Report.docx'
NAMES=dict(LABELS,sales_only='Sales permission only')
SELECT=['reference_100','bps_100','liquidity_cap','sales_only','anticipatory_12','anticipatory_sale']

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--largest',type=int,choices=[5000,10000],required=True);a=ap.parse_args()
    n=a.largest;directory=HERE/'outputs'
    result=json.loads((directory/f'stage_{n}_summary.json').read_text())
    verification=json.loads((HERE/'verification.json').read_text());assert verification['passed'] and verification['largest_n']==n
    perf=pd.read_csv(directory/f'stage_{n}_performance.csv');pairs=pd.read_csv(directory/f'stage_{n}_paired.csv')
    prim=pd.read_csv(directory/f'stage_{n}_prim_scores.csv');overlap=pd.read_csv(directory/f'stage_{n}_prim_overlap.csv')
    thresholds=pd.read_csv(directory/f'stage_{n}_prim_thresholds.csv');criteria=pd.read_csv(directory/f'stage_{n}_criteria.csv')
    lock=json.loads((HERE/'design_lock.json').read_text());sizes=sorted(perf.n.unique())
    original=pd.concat([pd.read_csv(ROOT/'outputs/all_main.csv'),pd.read_csv(ROOT/'outputs/diagnostic_sales.csv')])
    def values(policy,outcome='funded',size=n,horizon=str(CORE)):
        return perf[(perf.policy==policy)&(perf.outcome==outcome)&(perf.n==size)&(perf.horizon==horizon)].share
    def fmt(z,percent=True):
        z=np.asarray(z,float)*(100 if percent else 1)
        return f'{z.mean():.1f} ({z.min():.1f}–{z.max():.1f})'
    def effect(new,old,outcome='funded',size=n,horizon=str(CORE)):
        return pairs[(pairs.policy==new)&(pairs.comparator==old)&(pairs.outcome==outcome)&(pairs.n==size)&(pairs.horizon==horizon)].net_share
    # Use the existing report as the visual reference, preserving its styles.
    doc=Document(ROOT.parent/'deliverables/Strategy_RDM_Results_Report.docx')
    body=doc._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    doc.core_properties.title='Sampling convergence of the Strategy RDM analysis'
    doc.core_properties.subject='Independent Latin hypercube replications and PRIM stability'
    def para(text,style=None):return doc.add_paragraph(text,style)
    def head(text):doc.add_heading(text,level=1)
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        for j,h in enumerate(headers):t.rows[0].cells[j].text=h
        for row in rows:
            cells=t.add_row().cells
            for j,v in enumerate(row):cells[j].text=str(v)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if i==0:trpr.append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(row.cells):
                c.width=Inches(widths[j]);pr=c._tc.get_or_add_tcPr()
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'404040' if i==0 else ('F2F2F2' if i%2==0 else 'FFFFFF'));pr.append(fill)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'75');x.set(qn('w:type'),'dxa');margins.append(x)
                pr.append(margins)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size=Pt(9);run.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
        para('');return t
    policy_stable=result['criteria']['policy'] and result['criteria']['paired']
    prim_stable=result['criteria']['prim']
    para('Sampling convergence of the Strategy RDM analysis','Title')
    para('Independent Latin hypercube replications and PRIM stability','Subtitle')
    para('Research analysis for Murray A. Rudd  •  October 11 2026  •  Supplement to release 1.1.0')
    if policy_stable:
        opening=f'The funded shares and paired funded effects meet the prespecified convergence tolerances at {n:,} futures per replicate.'
    else:
        opening=f'Some funded shares or paired funded effects still exceed the prespecified convergence tolerances at {n:,} futures per replicate.'
    opening+=(' The primary PRIM regions also meet their stability criteria.' if prim_stable else ' The primary PRIM regions do not meet all stability criteria, so their exact boundaries should remain exploratory.')
    opening+=(' The substantive distinction between advance funding and Bitcoin sale permission holds in every new replicate at both long endpoints.' if result['criteria']['interpretation'] else ' At least one prespecified component interpretation does not hold in every replication; the detailed results identify the exception.')
    para(opening)
    para(f'The experiment evaluates {result["futures"]:,} new futures across {result["batches"]} independent Latin hypercube samples, with eleven policies per future: {result["evaluations"]:,} additional policy–future evaluations. It retains the model, uncertainty ranges, price generators, financing rules and corporate opening used in release 1.1. The variation below measures sensitivity to randomization within that exploratory design. It does not estimate probabilities of corporate outcomes.')
    head('Funded outcomes at five years')
    rows=[]
    for p in SELECT:
        old=original[(original.policy==p)&(original.horizon==str(CORE))].funded.mean()*100
        rows.append([NAMES[p],f'{old:.1f}']+[fmt(values(p,size=int(size))) for size in sizes])
    table(['Policy','Original 2,000']+[f'{int(s):,} per replicate' for s in sizes],rows,[2.15,.8]+[(6.8-2.95)/len(sizes)]*len(sizes))
    para('Percent meeting every payment and reserve requirement through June 2031. New designs show the mean followed by the minimum and maximum across three independent replications. Each replication compares policies on exactly matched futures. The original column is the retained single design; it is not included in the new replication means.','Caption')
    head('What the test establishes')
    pcrit=criteria[criteria.domain.isin(['policy','paired'])]
    para(f'Across all eleven policies and the six prespecified paired funded contrasts, the largest replicate range at {n:,} futures is {pcrit[pcrit.criterion=="replicate_range_pp"].value.max():.2f} percentage points. The largest mean change from the preceding sample size is {pcrit[pcrit.criterion=="mean_shift_pp"].value.max():.2f} points. The prospective limits were 2 and 1 points, respectively, at both June 2031 and June 2033.')
    para('A stable large policy advantage can coexist with uncertain ordering among policies separated by only a few cases. A narrow replicate range also cannot establish that the chosen uncertainty ranges or model structure capture all relevant futures.')

    doc.add_page_break();head('Policy effects across independent samples')
    rows=[]
    contrasts=[('sales_only','liquidity_cap','Add sales without advance funding'),('anticipatory_12','liquidity_cap','Add advance funding without sales'),('anticipatory_sale','sales_only','Add advance funding with sales'),('anticipatory_sale','anticipatory_12','Add sales with advance funding')]
    for new,old,label in contrasts:
        for outcome,label2 in [('payments','Payments'),('funded','Funded')]:
            rows.append([label,label2]+[fmt(effect(new,old,outcome,size=int(s))) for s in sizes])
    table(['Policy change','Target']+[f'{int(s):,} futures' for s in sizes],rows,[2.35,.65]+[(6.8-3.)/len(sizes)]*len(sizes))
    para('Net percentage-point gains at June 2031, expressed as the mean and replicate range. The package retains separate gains and losses; these effects are paired within each future. Every component comparison keeps the 12-month cash rule and preferred principal cap fixed.','Caption')
    para(f'At June 2033 in the largest samples, BPS 100 meets the funded target in {fmt(values("bps_100",horizon=str(LONG)))}% of futures, compared with {fmt(values("anticipatory_sale",horizon=str(LONG)))}% under combined advance funding and sale permission. These percentages retain the full design denominator.')
    # A compact black-and-white convergence figure, with individual seeds visible.
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False})
    fig,axes=plt.subplots(1,2,figsize=(7.2,3.0),sharex=True)
    for ax,h,title in zip(axes,[str(CORE),str(LONG)],['June 2031','June 2033']):
        for j,(p,marker,style,label) in enumerate([('bps_100','o','-','BPS 100'),('anticipatory_12','s','--','Advance funding'),('anticipatory_sale','^',':','Advance and sales')]):
            means=[]
            for i,size in enumerate(sizes):
                z=values(p,size=int(size),horizon=h).to_numpy()*100;means.append(z.mean())
                ax.scatter(i+np.linspace(-.055,.055,len(z))+(j-1)*.08,z,marker=marker,s=19,facecolors='white',edgecolors='black',linewidths=.75,zorder=3)
            ax.plot(np.arange(len(sizes))+(j-1)*.08,means,color='black',linestyle=style,linewidth=.9,label=label)
        ax.set_xticks(range(len(sizes)),[f'{int(s):,}' for s in sizes]);ax.set_xlabel('Futures per replicate');ax.set_title(title);ax.grid(axis='y',color='.88',linewidth=.5)
    axes[0].set_ylabel('Funded share (%)');axes[1].legend(frameon=False,fontsize=8,loc='center right')
    fig.tight_layout();(HERE/'figures').mkdir(exist_ok=True)
    for ext in ['png','svg']:fig.savefig(HERE/'figures'/f'policy_convergence.{ext}',dpi=220,bbox_inches='tight')
    plt.close(fig)
    p=para('');p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(HERE/'figures/policy_convergence.png'),width=Inches(6.65));shape._inline.docPr.set('descr','Funded shares across sample sizes, showing three independent seeds and their means for three policies at both long endpoints.')
    para('Figure 1. Individual markers show independent replications; lines connect their means. Horizontal offsets separate overlapping policies. The figure describes design sensitivity rather than a forecast distribution.','Caption')

    doc.add_page_break();head('Stability of the PRIM vulnerability regions')
    size_map={b['tag']:b['n'] for b in lock['batches']};prim['n']=prim.tag.map(size_map)
    rows=[]
    for outcome,label in [('funded','Funding'),('nominal_value','Value')]:
        for kind in ['first','union']:
            for size in sizes:
                z=prim[(prim.policy=='bps_100')&(prim.horizon==str(CORE))&(prim.outcome==outcome)&(prim.kind==kind)&(prim.split=='common_probe')&(prim.n==size)]
                jo=overlap[(overlap.policy=='bps_100')&(overlap.horizon==str(CORE))&(overlap.outcome==outcome)&(overlap.kind==kind)&(overlap.n==size)].mask_jaccard
                rows.append([label+' '+kind,f'{int(size):,}',fmt(z.density),fmt(z.coverage),f'{jo.min():.2f}'])
    table(['Failure target and region','Futures','Density %','Coverage %','Minimum overlap'],rows,[1.75,.7,1.65,1.65,1.05])
    para('New PRIM fits for BPS 100 at June 2031, evaluated on the same original 2,000 futures, which are excluded from every new fit. Density is the failure share inside the region; coverage is the share of all failures it captures. Parentheses show the range across replications. Overlap is the smallest pairwise Jaccard index for selected cases; 1 means identical regions.','Caption')
    para('First denotes the first selected box; union combines up to three boxes. Prospective limits are 10 percentage points for replicate ranges and adjacent-size mean shifts in density and coverage, minimum region overlap of 0.60, and minimum selected-variable/direction overlap of 0.50.','Caption')
    failed=criteria[(criteria.domain=='prim')&~criteria.passed]
    if len(failed):
        if set(failed.criterion)=={'mask_jaccard'} and set(failed.outcome)=={'funded'} and set(failed.kind)=={'first'}:
            observed=', '.join(f'{v:.3f}' for v in failed.value)
            para(f'The {len(failed)} remaining failures concern overlap of the first funding-failure box: Jaccard values are {observed}, below the 0.60 requirement. All primary density, coverage, and selected-variable/direction checks pass. The funding-box union and both nominal-value regions also pass their overlap checks. Thus the remaining instability is localized to which cases enter the first funding box; it does not invalidate every PRIM result.')
        else:
            failnames=', '.join(sorted(set(failed.criterion.str.replace('_',' '))))
            para(f'The final PRIM checks still fail for: {failnames}. A stable aggregate failure share does not ensure a unique explanatory box. Different combinations of related trajectory summaries can capture overlapping failure mechanisms. Exact cutoffs should remain exploratory.')
    else:para('The primary PRIM checks meet the locked tolerances. This supports the stability of these broad regions under repeated draws from the same design; it does not establish invariant financial thresholds or complete coverage of failure mechanisms.')
    hold=prim[(prim.policy=='bps_100')&(prim.horizon==str(CORE))&(prim.kind=='first')&(prim.split=='holdout')&(prim.n==n)]
    for o,label in [('funded','Funding'),('nominal_value','Nominal-value')]:
        z=hold[hold.outcome==o]
        para(f'{label} failure in the fresh holdouts has first-box density {fmt(z.density)}% and coverage {fmt(z.coverage)}%. These holdouts contain {n//4:,} futures per replicate and remain distinct from the common comparison set.')

    doc.add_page_break();head('Protocol and implications for the paper')
    stages=[]
    for size in [5000,10000]:
        path=directory/f'stage_{size}_summary.json'
        if not path.exists():continue
        stage=json.loads(path.read_text())
        stages.append([f'{size:,}',*['Pass' if stage['criteria'][k] else 'Not met' for k in ['policy','paired','interpretation','prim']]])
    table(['Largest sample','Funded shares','Paired effects','Component interpretation','PRIM'],stages,[1.15,1.2,1.2,2.05,1.2])
    para('The protocol was frozen before generating new outcomes. It requires three 10,000-future replications if any primary criterion fails after the 2,000-versus-5,000 comparison, and stops after that extension even if some criteria remain unmet. These are pragmatic tolerances selected for the study, not significance tests or universal sample-size standards.')
    if policy_stable:
        para(f'For the paper, 2,000 futures are sufficient to identify the broad policy finding, but one such sample does not establish numerical convergence. Funded policy comparisons pass the selected tolerances at 5,000 and remain stable at {n:,}. Retain the original design as the discovery analysis and use the larger replications to qualify its precision. Small differences between near-equal policies should not carry substantive weight.')
    else:
        para('For the paper, the central policy interpretation should be distinguished from numerical convergence. Report the remaining replicate variation and adjacent-size shifts directly. The protocol does not justify describing every percentage or policy ordering as fully stabilized.')
    para('PRIM should support the mechanism discussion at the level warranted by its observed stability. The supplied threshold tables retain every selected variable, inequality, cutoff and position within the common reference distribution. Commonly selected conditions can inform the interpretation even when alternate boxes or exact cutoffs remain unstable. The absence of a stable box is not evidence that the underlying financing mechanism is absent.')
    para('The checks do not change the economic scope. They add no automatic conversion, refinancing, preferred deferral, new price feedback or new historical data pool. They also do not repeat the original closure continuations or structural sensitivities at every larger sample size. Their conclusion concerns sampling of the existing exploratory design, conditional on its ranges, mixture and model relationships.')
    head('Reproducibility and verification')
    checks=sum(b['independent_checks'] for b in verification['batches'])
    para(f'Every one of the {result["evaluations"]:,} new policy–future evaluations reconciles its atomic ledger prefixes. The runs execute {checks:,} additional independent stock-flow, authority and entitlement checks. Verification confirms complete policy and checkpoint cross-products, unique future identifiers, exact Latin hypercube strata for all 15 continuous inputs, balanced binary margins, disjoint training and holdout partitions, and the recorded performance and PRIM validation counts.')
    para('The companion package retains the release 1.1 model and earlier evidence, plus a convergence directory containing the prospective lock, all generated inputs, completed outcome tables, paired gains and losses, PRIM fits and peeling trajectories, validation scores, overlap and threshold diagnostics, verification records, and the report builder. The original model and source commit remain unchanged. The supplement README gives the exact rerun commands.')
    para('Neither the number of policy evaluations nor the number of checkpoints increases the count of independently sampled futures. Repeated seeds test sensitivity to randomization within the chosen design; they do not confer empirical calibration or a probability interpretation on the resulting frequencies.')
    OUT.parent.mkdir(exist_ok=True);doc.save(OUT);print(str(OUT))

if __name__=='__main__':main()
